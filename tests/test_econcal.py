from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from econcal import load_events, news_blackout_mask, to_dataframe
from econcal.cli import main

SAMPLE = str(Path(__file__).with_name("sample-events.json"))


def test_dataframe():
    events = load_events(sample_file=SAMPLE, countries=["united states"], importance="high")
    df = to_dataframe(events)
    assert {"datetime_utc", "event", "actual", "forecast", "surprise"} <= set(df.columns)
    assert str(df["datetime_utc"].dt.tz) == "UTC"


def test_blackout_mask():
    events = load_events(sample_file=SAMPLE, countries=["united states"], importance="high")
    idx = pd.date_range("2026-10-14 11:00", "2026-10-14 14:00", freq="5min", tz="UTC")
    mask = news_blackout_mask(idx, events, before=15, after=30)
    assert mask[pd.Timestamp("2026-10-14 12:30", tz="UTC")]       # CPI release
    assert not mask[pd.Timestamp("2026-10-14 11:00", tz="UTC")]


def test_cli_xlsx(tmp_path):
    out = tmp_path / "cal.xlsx"
    assert main(["fetch", "--sample", SAMPLE, "--countries", "us,ez", "--format", "xlsx", "--out", str(out)]) == 0
    ws = load_workbook(out)["Calendar"]
    assert ws.max_row > 50 and ws["F1"].value in ("event", "importance", "currency") or ws.max_column > 8


def test_cli_blackout(tmp_path, capsys):
    out = tmp_path / "w.csv"
    assert main(["blackout", "--sample", SAMPLE, "--currencies", "USD", "--out", str(out)]) == 0
    lines = out.read_text().splitlines()
    assert lines[0].startswith("start_utc") and any("Nonfarm Payrolls" in l for l in lines)


def test_cli_next(capsys):
    assert main(["next", "--sample", SAMPLE, "--importance", "high", "--now", "2026-10-27T09:00:00Z", "--count", "5"]) == 0
    assert "Durable" in capsys.readouterr().out
